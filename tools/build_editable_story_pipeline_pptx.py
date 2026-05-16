#!/usr/bin/env python3
"""
Build an editable PowerPoint flowchart with native shapes, connectors, and text boxes.

Requires: python-pptx (unset proxy if needed: HTTP_PROXY= HTTPS_PROXY= pip install python-pptx)

Output: presentation/Auto_Story_Pipeline_editable.pptx
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt


def rgb(r: int, g: int, b: int) -> RGBColor:
    return RGBColor(r, g, b)


def style_rect(shape, fill: RGBColor | None, line: RGBColor | None, line_pt: float = 1.0) -> None:
    if fill is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()
    if line is not None:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_pt)


def put_text(
    shape,
    text: str,
    *,
    size: float = 10,
    align=PP_ALIGN.LEFT,
    bold: bool = False,
    color: RGBColor | None = None,
) -> None:
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.alignment = align
    p.space_after = Pt(0)
    if color is not None:
        p.font.color.rgb = color
    tf.margin_left = Inches(0.07)
    tf.margin_right = Inches(0.07)
    tf.margin_top = Inches(0.06)
    tf.margin_bottom = Inches(0.06)


def rounded(slide, x, y, w, h, text, *, fill, font=10, bold=False):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    put_text(sh, text, size=font, bold=bold)
    style_rect(sh, fill, rgb(100, 110, 125), 1)
    return sh


def label_only(slide, x, y, w, h, text, *, size=9, clr=RGBColor(45, 55, 70)):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    put_text(tb, text, size=size, align=PP_ALIGN.CENTER, color=clr)
    tb.fill.background()
    tb.line.fill.background()
    return tb


def connect(slide, x1, y1, x2, y2):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.width = Pt(1.85)
    c.line.color.rgb = rgb(72, 78, 95)
    return c


def output_bar(slide, y, yellow):
    tb = rounded(
        slide,
        0.18,
        y,
        12.98,
        0.74,
        "Output: N story panels (one per scene) · fully automated · no manual prompt or pixel edits per story after generation",
        fill=yellow,
        font=10.5,
        bold=True,
    )
    tb.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    style_rect(tb, yellow, rgb(160, 120, 40), 1.1)
    return tb


def build(out_path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])

    BLUE_BG = rgb(230, 240, 255)
    GREEN_BG = rgb(232, 246, 232)
    YELLOW_BG = rgb(255, 242, 204)
    PURPLE_BG = rgb(243, 232, 255)
    WHITE = rgb(255, 255, 255)

    # Z-order: large panels first (back)
    lp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.12), Inches(1.78), Inches(6.38), Inches(4.55))
    style_rect(lp, rgb(222, 235, 252), rgb(160, 190, 230), 0.55)

    rp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.64), Inches(1.78), Inches(6.38), Inches(4.55))
    style_rect(rp, rgb(222, 244, 222), rgb(160, 210, 160), 0.55)

    ttl = s.shapes.add_textbox(Inches(0.35), Inches(0.08), Inches(9.85), Inches(0.38))
    put_text(ttl, "Auto Story Pipeline — editable architecture", size=18, bold=True, align=PP_ALIGN.LEFT)

    # Notes top-right (no overlap with right column stacks)
    notes = rounded(
        s,
        10.52,
        0.12,
        2.62,
        1.52,
        "Notes\n• Consistent self-attention: OFF\n  in submitted runs\n• DiT / PixArt: exploratory\n  comparison only",
        fill=WHITE,
        font=8.5,
    )
    style_rect(notes, rgb(250, 250, 250), rgb(140, 140, 150), 0.8)

    rounded(
        s,
        0.22,
        0.48,
        3.05,
        1.05,
        "Scene-marked story text\n• [SCENE-n], [SEP]\n• <Entity> tags across scenes",
        fill=WHITE,
        font=9.5,
    )
    rounded(
        s,
        3.35,
        0.48,
        3.45,
        1.05,
        "Parser\n• Ordered scenes\n• Distinct tagged entities |E|\n• Raw text for planner",
        fill=PURPLE_BG,
        font=9.5,
    )
    router = s.shapes.add_shape(
        MSO_SHAPE.FLOWCHART_DECISION,
        Inches(7.4),
        Inches(0.44),
        Inches(1.78),
        Inches(1.12),
    )
    put_text(router, "Router\n|E|=1 vs |E|≥2", size=9, align=PP_ALIGN.CENTER, bold=True)
    style_rect(router, rgb(255, 248, 220), rgb(130, 100, 60), 1.1)

    connect(s, 3.27, 1.0, 3.35, 1.0)
    connect(s, 6.8, 1.0, 7.4, 0.92)

    label_only(s, 0.25, 1.84, 6.1, 0.26, "Single-character path (submitted when |E| = 1)", size=10, clr=rgb(20, 55, 110))
    label_only(s, 6.75, 1.84, 5.8, 0.26, "Multi-character path (|E| ≥ 2)", size=10, clr=rgb(15, 90, 35))

    ix, iw, iy0 = 0.28, 5.92, 2.22
    gap = 0.07
    heights = [0.66, 0.66, 0.82, 0.5, 0.68, 0.52]
    texts = [
        "LLM-assisted structured planner (text only)\n• Character specs, scene fields, continuity, setting focus",
        "Subject-type normalization + scene-consistency phrases\n• Humans / animals / robots",
        "Anchor Bank → canonical half-body anchor\n(portrait candidates → one reference for conditioning)",
        "Identity gate · single clear target per scene",
        "SDXL-Turbo + IP-Adapter · identity strength λ = 0.3",
        "Optional: CLIP ViT-B/32 panel selection · multi-candidate runs",
    ]
    ys = iy0
    left_shapes: list[tuple[float, float]] = []
    for h, txt in zip(heights, texts):
        rounded(s, ix, ys, iw, h, txt, fill=BLUE_BG, font=8.75)
        left_shapes.append((ys, h))
        ys += h + gap

    rx, rw, ry0 = 6.92, 5.92, 2.22
    rheights = [0.72, 0.86, 0.88]
    rtexts = [
        "LLM planner + StoryDiffusion natural prompt adapter",
        "Identity bank inputs · [Character] lines · identity ref prompts · natural scene prompts",
        "Native StoryDiffusion · no forced single-anchor IP-Adapter on two-primary scenes",
    ]
    ry = ry0
    for h, txt in zip(rheights, rtexts):
        rounded(s, rx, ry, rw, h, txt, fill=GREEN_BG, font=9)
        ry += h + gap

    # Branch arrows from router
    connect(s, 8.3, 1.38, 3.52, 2.05)
    connect(s, 8.95, 1.38, 10.88, 2.05)
    label_only(s, 5.25, 1.22, 0.82, 0.28, "|E| = 1", size=8.8)
    label_only(s, 9.6, 1.22, 0.9, 0.28, "|E| ≥ 2", size=8.8)

    cx = ix + iw / 2
    for i in range(len(left_shapes) - 1):
        connect(s, cx, left_shapes[i][0] + left_shapes[i][1], cx, left_shapes[i + 1][0])

    rcx = rx + rw / 2
    rstack_y = [ry0, ry0 + rheights[0] + gap, ry0 + rheights[0] + gap + rheights[1] + gap]
    for i in range(2):
        connect(s, rcx, rstack_y[i] + rheights[i], rcx, rstack_y[i + 1])

    out_y = 6.45
    output_bar(s, out_y, YELLOW_BG)

    connect(s, cx, left_shapes[-1][0] + left_shapes[-1][1], cx, out_y)
    connect(s, rcx, rstack_y[2] + rheights[2], rcx, out_y)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    print("Saved", out_path)


if __name__ == "__main__":
    repo = Path(__file__).resolve().parents[1]
    build(repo / "presentation" / "Auto_Story_Pipeline_editable.pptx")
